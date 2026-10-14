import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from malote import anexos, controle, falso
from malote.registro import Config, Registro


class Anexos(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_cnpjs_em_xml_txt_e_zip(self):
        xml = self.pasta / "n.xml"
        xml.write_bytes(falso.nota_xml(1, "ti", "101", 10.0, "x", (2026, 9)))
        achados = anexos.cnpjs(xml, "n.xml")
        self.assertIn(falso.cnpj_da_empresa("101"), achados)
        self.assertIn(falso.cnpj_ficticio(903), achados)
        txt = self.pasta / "n.txt"
        txt.write_text("Tomador CNPJ 12.345.678/0001-95 e 12345678000195 e 1234567800019", encoding="utf-8")
        self.assertEqual(anexos.cnpjs(txt, "n.txt"), {"12345678000195"})
        z = self.pasta / "n.zip"
        z.write_bytes(falso.zip_de({"a.xml": xml.read_bytes(), "b.txt": b"nada"}))
        self.assertIn(falso.cnpj_da_empresa("101"), anexos.cnpjs(z, "n.zip"))
        self.assertEqual(anexos.cnpjs(txt, "n.docx"), set())
        ruim = self.pasta / "ruim.zip"
        ruim.write_bytes(b"isto nao e zip")
        self.assertEqual(anexos.cnpjs(ruim, "ruim.zip"), set())

    @unittest.skipUnless(importlib.util.find_spec("pdfminer"), "pdfminer.six não instalado")
    def test_cnpj_em_pdf(self):
        pdf = self.pasta / "n.pdf"
        pdf.write_bytes(falso.nota_pdf(5, "limpeza", "102", 99.0, "limpeza", (2026, 9)))
        self.assertIn(falso.cnpj_da_empresa("102"), anexos.cnpjs(pdf, "n.pdf"))

    def test_qual_empresa(self):
        cand = {"101 - A": ["11111111000191"], "102 - B": ["22222222000191", ""], "103 - C": ["33333333000191"]}
        self.assertEqual(anexos.qual_empresa({"22222222000191"}, cand), "102 - B")
        self.assertIsNone(anexos.qual_empresa({"11111111000191", "22222222000191"}, cand))
        self.assertIsNone(anexos.qual_empresa(set(), cand))
        # filial: mesma raiz, outro final
        self.assertEqual(anexos.qual_empresa({"33333333000272"}, cand), "103 - C")


class Controle(unittest.TestCase):
    def test_planilhas(self):
        with tempfile.TemporaryDirectory() as tmp:
            arquivo = Path(tmp) / "sub" / "Tomados 09-2026.xlsx"
            controle.gerar(arquivo, [["01/10/2026 09:00", "101 - A", "Maria", "m@a.example", "Mandou anexo", 1, "", "", "Segue"]],
                           [["102 - B", "Ana", "b@b.example"]], [["02/10/2026 14:00", "x@y.example", ""]])
            wb = load_workbook(arquivo)
            self.assertEqual(wb.sheetnames, ["Respostas", "Sem resposta", "Não entregues"])
            self.assertEqual(wb["Respostas"].max_row, 2)
            self.assertEqual(wb["Respostas"]["A2"].fill.fgColor.rgb[-6:], controle.CORES["Mandou anexo"])
            conf = Path(tmp) / "Conferência.xlsx"
            controle.gerar_conferencia(conf, [["103", "C", "Ana"]], [], [["c@x.example", 2, "A; B"]], [], [], [], [])
            wb = load_workbook(conf)
            self.assertEqual(len(wb.sheetnames), 7)
            self.assertEqual(wb["Sem e-mail"]["A2"].value, "103")


class RegistroEConfig(unittest.TestCase):
    def test_registro(self):
        with tempfile.TemporaryDirectory() as tmp:
            pasta = Path(tmp) / "dados"
            r = Registro(pasta)
            r.aprender("X@Y.example", "101 - A")
            r.marcar_salvo("id1", "101 - A", "09/2026", r"C:\x", ["a.pdf"])
            r.marcar_respondido("id1", "x@y.example", "Ok")
            r.marcar_reenviado(["A@B.example"], "09/2026")
            de_novo = Registro(pasta)
            self.assertEqual(de_novo.remetentes, {"x@y.example": "101 - A"})
            self.assertEqual(de_novo.salvos["id1"]["arquivos"], ["a.pdf"])
            self.assertEqual(de_novo.respondidos["id1"]["para"], "x@y.example")
            self.assertTrue(de_novo.reenviados["a@b.example"].startswith("09/2026 em "))
            c = Config(pasta)
            self.assertIsNone(c.get("conta"))
            c.set("conta", "f@c.example")
            self.assertEqual(Config(pasta).get("conta"), "f@c.example")
            self.assertEqual(json.loads((pasta / "config.json").read_text(encoding="utf-8")), {"conta": "f@c.example"})


if __name__ == "__main__":
    unittest.main()
