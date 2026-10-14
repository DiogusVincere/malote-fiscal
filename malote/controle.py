"""Planilhas de controle: quem respondeu, o que disse, onde foi parar, quem falta cobrar; e a conferência
da planilha de clientes (o que vale corrigir na origem).

Gravadas na pasta de dados da ferramenta (`controle\\`), nunca na pasta do cliente.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import tema

CABECALHO = tema.EXCEL_CABECALHO_FUNDO[2:]
CORES = {"Mandou anexo": tema.EXCEL_VERDE_T[2:], "Mandou link": tema.EXCEL_VERDE_T[2:],
         "Sem serviços": tema.EXCEL_ZEBRA[2:], "Ler": tema.EXCEL_AMBAR_T[2:],
         "Não entregue": tema.EXCEL_VERMELHO_T[2:], "Sem resposta": tema.EXCEL_AMBAR_T[2:]}


def _aba(wb: Workbook, titulo: str, cabecalho: list[str], linhas: list[list], larguras: list[int], cor_por=None):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CABECALHO)
        c.alignment = Alignment(vertical="center")
    for linha in linhas:
        ws.append(linha)
        if cor_por is not None and linha[cor_por] in CORES:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=CORES[linha[cor_por]])
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def gerar(arquivo: Path, respostas: list[list], sem_resposta: list[list], nao_entregues: list[list]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _aba(wb, "Respostas",
         ["Recebido", "Cliente", "Remetente", "E-mail", "Situação", "Anexos", "Salvo em", "Respondido", "O que disse"],
         respostas, [16, 42, 26, 32, 14, 8, 60, 12, 80], cor_por=4)
    _aba(wb, "Sem resposta", ["Empresa", "Resp. fiscal", "E-mails perguntados"], sem_resposta, [50, 28, 70])
    _aba(wb, "Não entregues", ["Recebido", "E-mail que voltou", "Cliente"], nao_entregues, [16, 40, 42])
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    wb.save(arquivo)
    return arquivo


def gerar_conferencia(arquivo: Path, sem_email: list[list], internos: list[list], compartilhados: list[list],
                      sem_pasta: list[list], fora_da_lista: list[list], so_grupo: list[list],
                      so_planilha: list[list]) -> Path:
    """O que vale corrigir na origem (planilha de clientes / grupos do Outlook)."""
    wb = Workbook()
    wb.remove(wb.active)
    _aba(wb, "Só nos grupos", ["E-mail", "Grupo(s)", "Observação"], so_grupo, [40, 40, 50])
    _aba(wb, "Só na planilha", ["E-mail", "Empresa(s)"], so_planilha, [40, 100])
    _aba(wb, "Sem e-mail", ["Código", "Razão social", "Resp. fiscal"], sem_email, [16, 60, 30])
    _aba(wb, "E-mail interno", ["E-mail", "Empresas"], internos, [40, 90])
    _aba(wb, "E-mail em várias empresas", ["E-mail", "Qtde", "Empresas"], compartilhados, [40, 8, 120])
    _aba(wb, "Código sem pasta", ["Código", "Razão social", "Resp. fiscal"], sem_pasta, [16, 60, 30])
    _aba(wb, "Não recebem", ["E-mail", "De onde", "Motivo"], fora_da_lista, [40, 18, 50])
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    wb.save(arquivo)
    return arquivo
