"""Planilha de clientes ("Clientes AAAA-MM-DD.xlsx"): a fonte de quem recebe e de quem é cada e-mail.

Colunas esperadas (a ordem não importa; o título é procurado sem acento): Código, Razão Social,
CNPJ/CPF, E-mails, Regime municipal, Resp. fiscal e, opcionalmente, Ex-cliente. É o formato que
os sistemas de gestão de escritório costumam exportar.

Regras:
- a lista é a planilha JUNTO com os grupos do Outlook que a pessoa marcar. Os grupos e contatos
  do Outlook nunca são alterados;
- da planilha recebe quem não está marcado como "ISENTA/INEXISTENTE MUNICIPAL" (regime em branco
  recebe, porque cadastro incompleto é comum e empresa em branco manda nota). Ex-cliente fica de fora;
- e-mail que aparece em várias empresas recebe uma mensagem só.

O que uma exportação real costuma ter de estranho (tratado aqui):
- coluna "E-mails" com vários endereços separados por " | ", "*" na frente de um endereço e
  e-mails da própria equipe no meio dos clientes (esses não recebem a pergunta);
- o mesmo e-mail em várias empresas (contador de um grupo, dono de várias lojas);
- código que não é número: "000 EX-CLIENTE", CNPJ usado como código, apelido.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .competencia import sem_acento

# Só fica de fora quem o cadastro diz que NÃO tem tomados. Regime em branco recebe.
REGIMES_SEM_TOMADOS = {"ISENTA MUNICIPAL", "INEXISTENTE MUNICIPAL"}
_EMAIL = re.compile(r"^[\w.+-]+@[\w-]+(\.[\w-]+)+$")


@dataclass
class Empresa:
    codigo: str
    razao: str
    cnpj: str                   # só dígitos
    emails: list[str] = field(default_factory=list)
    regime_municipal: str = ""
    resp_fiscal: str = ""
    ex_cliente: bool = False

    @property
    def recebe(self) -> bool:
        return not self.ex_cliente and self.regime_municipal.upper().strip() not in REGIMES_SEM_TOMADOS

    @property
    def rotulo(self) -> str:
        return f"{self.codigo} - {self.razao}"


def _chave(titulo) -> str:
    return re.sub(r"[^a-z0-9]+", " ", sem_acento(str(titulo or "")).lower()).strip()


def emails_do_texto(texto: str) -> list[str]:
    achados = []
    for pedaco in re.split(r"[|,;\s]+", texto or ""):
        e = pedaco.strip().strip("*'\"<>()").lower()
        if _EMAIL.match(e) and e not in achados:
            achados.append(e)
    return achados


def achar(pastas: list[Path], prefixo: str = "Clientes") -> Path | None:
    """A planilha mais nova cujo nome começa com `prefixo`: pela data no nome ("Clientes 2026-10-01.xlsx"),
    senão pela data do arquivo."""
    candidatos = []
    for pasta in pastas:
        if pasta and Path(pasta).is_dir():
            candidatos += [p for p in Path(pasta).glob(f"{prefixo}*.xlsx") if not p.name.startswith("~$")]

    def data(p: Path):
        m = re.search(r"(\d{4})-(\d{2})-(\d{2})", p.name)
        return (datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m
                else datetime.fromtimestamp(p.stat().st_mtime))
    return max(candidatos, key=data) if candidatos else None


def ler(caminho: Path) -> list[Empresa]:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()  # o openpyxl confia na dimensão declarada e pode cortar linhas
        linhas = ws.iter_rows(values_only=True)
        titulos = [_chave(t) for t in next(linhas)]

        def coluna(*nomes: str) -> int | None:
            for nome in nomes:
                for i, t in enumerate(titulos):
                    if t == nome or t.endswith(" " + nome):
                        return i
            return None

        c_cod, c_razao, c_cnpj = coluna("codigo"), coluna("razao social"), coluna("cnpj cpf", "cnpj")
        c_emails, c_regime = coluna("e mails", "emails"), coluna("regime municipal")
        c_fiscal, c_ex = coluna("resp fiscal"), coluna("ex cliente")
        faltando = [n for n, c in (("Código", c_cod), ("Razão Social", c_razao), ("E-mails", c_emails),
                                   ("Regime municipal", c_regime)) if c is None]
        if faltando:
            raise ValueError(f"A planilha {caminho.name} não tem a(s) coluna(s): {', '.join(faltando)}.")

        def valor(linha, c):
            return "" if c is None or c >= len(linha) or linha[c] is None else str(linha[c]).strip()

        empresas = []
        for linha in linhas:
            codigo = valor(linha, c_cod)
            if not codigo:
                continue
            empresas.append(Empresa(
                codigo=codigo, razao=valor(linha, c_razao), cnpj=re.sub(r"\D", "", valor(linha, c_cnpj)),
                emails=emails_do_texto(valor(linha, c_emails)), regime_municipal=valor(linha, c_regime),
                resp_fiscal=valor(linha, c_fiscal),
                ex_cliente="EX-CLIENTE" in codigo.upper() or valor(linha, c_ex).lower() == "sim"))
        return empresas
    finally:
        wb.close()


def por_email(empresas: list[Empresa], so_ativas: bool = True) -> dict[str, list[Empresa]]:
    mapa: dict[str, list[Empresa]] = {}
    for e in empresas:
        if so_ativas and e.ex_cliente:
            continue
        for email in e.emails:
            mapa.setdefault(email, []).append(e)
    return mapa


def _interno(email: str, dominio_interno: str) -> bool:
    return bool(dominio_interno) and email.lower().endswith("@" + dominio_interno.lower())


def _quem_fica_de_fora(empresas: list[Empresa], fora_da_lista: dict[str, str], manuais: dict[str, list[Empresa]],
                       respondentes: set[str], principais: set[str], dominio_interno: str = "") -> dict[str, str]:
    """Empresa com vários e-mails: recebe o principal fixado; senão quem respondeu; senão todos.

    Devolve {e-mail que fica de fora: por quê}. Um e-mail de várias empresas (contador, dono de
    grupo) só fica de fora se não for o escolhido de nenhuma delas.
    """
    da_empresa: dict[str, list[str]] = {}
    for e in empresas:
        if e.recebe:
            da_empresa.setdefault(e.codigo, [])
            da_empresa[e.codigo] += [m for m in e.emails if m not in da_empresa[e.codigo]]
    for email, lista in manuais.items():
        for e in lista:
            if email not in da_empresa.setdefault(e.codigo, []):
                da_empresa[e.codigo].append(email)
    escolhidos: set[str] = set()
    fora: dict[str, str] = {}
    for codigo, emails in da_empresa.items():
        emails = [m for m in emails if m not in fora_da_lista and not _interno(m, dominio_interno)]
        if len(emails) <= 1:
            escolhidos.update(emails)
            continue
        fixos = [m for m in emails if m in principais]
        responderam = [m for m in emails if m in respondentes]
        vai = fixos or responderam or emails
        escolhidos.update(vai)
        motivo = ("principal da empresa: " if fixos else "a empresa responde por ") + ", ".join(vai)
        for m in emails:
            if m not in vai:
                fora.setdefault(m, f"{motivo} ({codigo})")
    return {m: motivo for m, motivo in fora.items() if m not in escolhidos}


@dataclass
class Destinatario:
    email: str
    empresas: list[Empresa]
    grupos: list[str]
    incluido: bool
    observacao: str = ""
    padrao: bool = True     # se receberia sem o usuário mexer (para saber o que guardar ao clicar)
    manual: bool = False    # acrescentado à mão no cadastro

    @property
    def origem(self) -> str:
        na_planilha = any(e.recebe for e in self.empresas) and not self.manual
        if na_planilha and self.grupos:
            return "planilha + grupo"
        if na_planilha:
            return "só planilha"
        return "só grupo" if self.grupos else "manual"


@dataclass
class ListaDeEnvio:
    destinatarios: list[Destinatario]
    lotes: list[list[str]]
    empresas: int                 # empresas da planilha que recebem
    sem_email: list[Empresa]      # recebem, mas não têm e-mail utilizável
    repetidos: int                # quantas vezes um e-mail serviria a outra empresa (vai uma vez só)


def lista_de_envio(empresas: list[Empresa], fora_da_lista: dict[str, str], tamanho: int = 40,
                   grupos: dict[str, list[str]] | None = None, incluir: set[str] | None = None,
                   manuais: dict[str, list[Empresa]] | None = None, respondentes: set[str] | None = None,
                   principais: set[str] | None = None, dominio_interno: str = "") -> ListaDeEnvio:
    """Junta planilha, grupos do Outlook e o que foi acrescentado à mão, sem repetir e-mail.

    grupos: {e-mail: [grupos em que está]}. incluir: marcados à mão mesmo sendo de ex-cliente.
    fora_da_lista: desmarcados, pediram para sair, trocaram de e-mail ou voltam sempre.
    manuais: {e-mail: [empresas]} acrescentados no cadastro.
    dominio_interno: domínio da conta que envia; e-mail desse domínio é da própria equipe e não recebe.
    """
    grupos, incluir, manuais = grupos or {}, incluir or set(), manuais or {}
    preteridos = _quem_fica_de_fora(empresas, fora_da_lista, manuais, respondentes or set(), principais or set(),
                                    dominio_interno)
    mapa = por_email(empresas, so_ativas=False)
    sem_email, repetidos, n = [], 0, 0
    da_planilha: set[str] = set()
    internos: set[str] = set()   # e-mails da própria equipe que vieram na planilha: aparecem, mas não recebem
    for e in empresas:
        if not e.recebe:
            continue
        n += 1
        uteis = [m for m in e.emails if not _interno(m, dominio_interno)]
        internos |= {m for m in e.emails if _interno(m, dominio_interno)}
        if not uteis:
            sem_email.append(e)
        for m in uteis:
            repetidos += m in da_planilha
            da_planilha.add(m)
    candidatos = da_planilha | set(grupos) | set(manuais) | internos
    destinatarios = []
    for email in sorted(candidatos):
        dono = mapa.get(email, [])
        manual = email in manuais and email not in da_planilha
        if manual:
            dono = manuais[email] or dono
        d = Destinatario(email, dono, sorted(grupos.get(email, [])), True, manual=manual)
        if manual:
            d.observacao = "acrescentado à mão"
        elif _interno(email, dominio_interno):
            d.padrao, d.observacao = False, "e-mail interno"
        elif not any(x.recebe for x in dono):
            if dono and all(x.ex_cliente for x in dono):
                d.padrao, d.observacao = False, "ex-cliente na planilha"
            elif dono:
                d.observacao = "regime municipal: " + ", ".join(sorted({x.regime_municipal for x in dono}))
            else:
                d.observacao = "não está na planilha"
        elif not d.grupos and grupos:
            d.observacao = "não está nos grupos"
        if email in preteridos and d.padrao:
            d.padrao, d.observacao = False, preteridos[email]
        d.incluido = d.padrao
        if email in fora_da_lista:
            d.incluido, d.observacao = False, fora_da_lista[email]
        elif email in incluir:
            d.incluido = True
        destinatarios.append(d)
    vistos = [d.email for d in destinatarios if d.incluido]
    lotes = [vistos[i:i + tamanho] for i in range(0, len(vistos), tamanho)]
    return ListaDeEnvio(destinatarios, lotes, n, sem_email, repetidos)
